"""Secure, configuration-bound participant and submission imports.

Uploaded bytes and raw external references remain transient. Durable matching
uses a keyed digest supplied through :class:`IdentityProtector`.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from collections import Counter
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from enum import StrEnum
from pathlib import PurePath
from typing import Any

from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.worksheet.datavalidation import (  # type: ignore[import-untyped]
    DataValidation,
)

from poli_insight.application.participation_policy import consent_policy
from poli_insight.application.ports.identity_protection import IdentityProtector
from poli_insight.application.ports.operations_repository import (
    AdminImportConsentDisposition,
    ImportedIdentityAuthority,
    ParticipantImportReference,
    ParticipantSubmissionImportBatch,
    ParticipantSubmissionImportRecord,
)
from poli_insight.application.ports.unit_of_work import UnitOfWork
from poli_insight.application.use_cases._operational_audit import (
    operational_audit_event,
)
from poli_insight.core.ids import new_id
from poli_insight.core.time import utc_now
from poli_insight.domain.content_hash import hash_json, sha256_digest
from poli_insight.domain.enum import (
    ActorType,
    AuditAction,
    QuestionType,
    ResponseFormat,
    SessionStatus,
    SubmissionStatus,
)
from poli_insight.domain.participation import (
    Participant,
    ParticipantAccessGrant,
    ParticipantIdentity,
)
from poli_insight.domain.scenario import ScenarioSnapshot
from poli_insight.domain.session import (
    ResponseQuestionDefinition,
    Session,
    SessionConfigurationVersion,
    SessionStakeholderGroup,
)
from poli_insight.domain.submission import Submission, SubmissionAnswer
from poli_insight.infrastructure.auth.tokens import generate_token
from poli_insight.infrastructure.security.identity_protection import (
    normalize_email_for_lookup,
    normalize_participant_reference,
)

TEMPLATE_SCHEMA_VERSION = "1"
SUPPORTED_EXTENSIONS = frozenset({".csv", ".xlsx"})
RESERVED_COLUMNS = (
    "template_schema_version",
    "session_slug",
    "configuration_version",
    "participant_ref",
    "participant_alias",
    "participant_name",
    "participant_email",
    "group",
    "record_state",
    "recorded_at",
)
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
_MAX_XLSX_MEMBER_BYTES = 25 * 1024 * 1024
_SAFE_FILENAME = re.compile(r"[^A-Za-z0-9._-]+")

UnitOfWorkFactory = Callable[[], UnitOfWork]
Clock = Callable[[], datetime]
IdFactory = Callable[[], str]


class ParticipantSubmissionImportError(ValueError):
    """Safe expected failure at the import trust boundary."""


class ImportRowStatus(StrEnum):
    READY = "ready"
    SKIPPED_CONFLICT = "skipped_conflict"
    PLANNED_REPLACEMENT = "planned_replacement"
    BLOCKING_ERROR = "blocking_error"


class ImportRecordState(StrEnum):
    ENROLLED = "enrolled"
    DRAFT = "draft"
    SUBMITTED = "submitted"


@dataclass(frozen=True, slots=True)
class ImportIssue:
    row_number: int
    column: str
    code: str
    message: str
    accepted_values: tuple[str, ...] = ()
    blocking: bool = True


@dataclass(frozen=True, slots=True, repr=False)
class ParticipantSubmissionImportRow:
    row_number: int
    participant_ref: str
    participant_alias: str
    group_key: str | None
    group_name: str | None
    record_state: ImportRecordState
    recorded_at: datetime | None
    answer_count: int
    has_identity: bool
    status: ImportRowStatus
    existing_participant_id: str | None
    issues: tuple[ImportIssue, ...]

    def __repr__(self) -> str:
        return (
            "ParticipantSubmissionImportRow("
            f"row_number={self.row_number}, participant_ref=<redacted>, "
            f"status={self.status.value!r}, issue_count={len(self.issues)})"
        )


@dataclass(frozen=True, slots=True)
class ImportPreviewSummary:
    total_rows: int
    new_participants: int
    enrolled_only: int
    drafts: int
    submitted: int
    skipped_conflicts: int
    planned_replacements: int
    identity_rows: int
    warning_count: int
    blocking_error_count: int


@dataclass(frozen=True, slots=True, repr=False)
class ParticipantSubmissionImportPreview:
    session_id: str
    scenario_snapshot_id: str
    configuration_version_id: str
    filename: str
    file_hash: str
    plan_hash: str
    rows: tuple[ParticipantSubmissionImportRow, ...]
    summary: ImportPreviewSummary
    error_report_csv: bytes = field(repr=False)

    @property
    def has_blocking_errors(self) -> bool:
        return self.summary.blocking_error_count > 0


@dataclass(frozen=True, slots=True, repr=False)
class PreviewParticipantSubmissionImportCommand:
    session_id: str
    filename: str
    content: bytes = field(repr=False)
    replace_row_numbers: frozenset[int] = frozenset()

    def __repr__(self) -> str:
        return (
            "PreviewParticipantSubmissionImportCommand("
            f"session_id={self.session_id!r}, filename={self.filename!r}, "
            "content=<redacted>)"
        )


@dataclass(frozen=True, slots=True)
class GenerateParticipantImportTemplateCommand:
    session_id: str
    file_format: str
    filled_example: bool = False


@dataclass(frozen=True, slots=True, repr=False)
class GeneratedParticipantImportTemplate:
    filename: str
    media_type: str
    content: bytes = field(repr=False)
    question_count: int

    def __repr__(self) -> str:
        return (
            "GeneratedParticipantImportTemplate("
            f"filename={self.filename!r}, content=<redacted>, "
            f"question_count={self.question_count})"
        )


@dataclass(frozen=True, slots=True)
class ParticipantImportSessionSummary:
    session_id: str
    session_title: str
    session_slug: str
    scenario_title: str
    scenario_version: str
    configuration_version: int
    lifecycle_state: SessionStatus
    response_format: ResponseFormat
    response_target_type: str
    scale_name: str
    scale_version: int
    groups: tuple[tuple[str, str], ...]
    required_question_count: int
    allow_incomplete_submission: bool
    allow_resubmissions: bool
    consistency_threshold: str | None
    consent_policy: str
    identity_policy: str


class GetParticipantImportSession:
    def __init__(self, unit_of_work_factory: UnitOfWorkFactory) -> None:
        self._unit_of_work_factory = unit_of_work_factory

    def execute(self, session_id: str) -> ParticipantImportSessionSummary:
        with self._unit_of_work_factory() as unit_of_work:
            scope = _load_scope(unit_of_work, session_id)
        scale = next(
            item for item in scope.snapshot.scales
            if item.scale_id == scope.configuration.scale_id
        )
        consent = consent_policy(scope.configuration)
        return ParticipantImportSessionSummary(
            session_id=scope.session.session_id,
            session_title=scope.session.title,
            session_slug=scope.session.public_slug,
            scenario_title=scope.snapshot.title,
            scenario_version=scope.snapshot.declared_version,
            configuration_version=scope.configuration.version_number,
            lifecycle_state=scope.session.status,
            response_format=scope.configuration.response_format,
            response_target_type=scope.configuration.response_target_type.value,
            scale_name=scale.name,
            scale_version=scale.definition_version,
            groups=tuple(
                (group.group_key, group.name)
                for group in sorted(
                    scope.configuration.stakeholder_groups,
                    key=lambda item: (item.display_order, item.group_key),
                )
                if group.is_active
            ),
            required_question_count=sum(
                question.required
                for question in scope.configuration.question_definitions
            ),
            allow_incomplete_submission=(
                scope.configuration.allow_incomplete_submission
            ),
            allow_resubmissions=scope.configuration.allow_resubmissions,
            consistency_threshold=(
                None
                if scope.configuration.consistency_threshold is None
                else str(scope.configuration.consistency_threshold)
            ),
            consent_policy=(
                f"Required · version {consent.version}"
                if consent.required
                else "Not required"
            ),
            identity_policy=scope.session.identity_policy,
        )


@dataclass(frozen=True, slots=True, repr=False)
class _ParsedRow:
    row_number: int
    participant_ref: str
    reference_digest: str | None
    participant_alias: str
    participant_name: str | None
    participant_email: str | None
    group: SessionStakeholderGroup | None
    record_state: ImportRecordState
    recorded_at: datetime | None
    answers: tuple[tuple[ResponseQuestionDefinition, str], ...]
    issues: tuple[ImportIssue, ...]


@dataclass(frozen=True, slots=True)
class _ImportScope:
    session: Session
    configuration: SessionConfigurationVersion
    snapshot: ScenarioSnapshot
    scale_values_by_input: Mapping[str, tuple[str, ...]]
    scale_choices: tuple[str, ...]
    question_headers: Mapping[str, ResponseQuestionDefinition]
    groups_by_input: Mapping[str, tuple[SessionStakeholderGroup, ...]]


class GenerateParticipantImportTemplate:
    def __init__(self, unit_of_work_factory: UnitOfWorkFactory) -> None:
        self._unit_of_work_factory = unit_of_work_factory

    def execute(
        self, command: GenerateParticipantImportTemplateCommand
    ) -> GeneratedParticipantImportTemplate:
        requested_format = command.file_format.strip().lower().removeprefix(".")
        if requested_format not in {"csv", "xlsx"}:
            raise ParticipantSubmissionImportError(
                "Templates are available only as UTF-8 CSV or .xlsx workbooks."
            )
        with self._unit_of_work_factory() as unit_of_work:
            scope = _load_scope(unit_of_work, command.session_id)
        headers = (*RESERVED_COLUMNS, *scope.question_headers)
        example = _example_row(scope) if command.filled_example else None
        suffix = "example" if command.filled_example else "blank"
        base = _safe_filename(f"{scope.session.public_slug}-{suffix}-import")
        if requested_format == "csv":
            content = _csv_template(headers, example)
            media_type = "text/csv"
        else:
            content = _xlsx_template(scope, headers, example)
            media_type = (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        return GeneratedParticipantImportTemplate(
            filename=f"{base}.{requested_format}",
            media_type=media_type,
            content=content,
            question_count=len(scope.question_headers),
        )


class PreviewParticipantSubmissionImport:
    def __init__(
        self,
        unit_of_work_factory: UnitOfWorkFactory,
        identity_protector: IdentityProtector,
        *,
        max_bytes: int,
        max_rows: int,
        clock: Clock = utc_now,
    ) -> None:
        self._unit_of_work_factory = unit_of_work_factory
        self._identity_protector = identity_protector
        self._max_bytes = max_bytes
        self._max_rows = max_rows
        self._clock = clock

    def execute(
        self, command: PreviewParticipantSubmissionImportCommand
    ) -> ParticipantSubmissionImportPreview:
        filename = _safe_uploaded_filename(command.filename)
        _validate_upload(filename, command.content, max_bytes=self._max_bytes)
        now = self._clock()
        with self._unit_of_work_factory() as unit_of_work:
            scope = _load_scope(unit_of_work, command.session_id)
            source_rows = _read_upload(
                filename,
                command.content,
                scope=scope,
                max_rows=self._max_rows,
            )
            parsed = tuple(
                _parse_source_row(
                    row_number,
                    source,
                    scope=scope,
                    protector=self._identity_protector,
                    now=now,
                )
                for row_number, source in source_rows
            )
            counts = Counter(
                row.reference_digest
                for row in parsed
                if row.reference_digest is not None
            )
            existing = unit_of_work.participant_imports.participants_by_reference_digests(
                command.session_id,
                tuple(counts),
            )
            participant_by_id = {
                participant.participant_id: participant
                for participant in unit_of_work.participants.get_many(
                    tuple(existing.values())
                )
            }

        classified = tuple(
            _classify_row(
                row,
                scope=scope,
                duplicate=(
                    row.reference_digest is not None
                    and counts[row.reference_digest] > 1
                ),
                existing_participant_id=(
                    None
                    if row.reference_digest is None
                    else existing.get(row.reference_digest)
                ),
                participant_by_id=participant_by_id,
                replace_selected=row.row_number in command.replace_row_numbers,
            )
            for row in parsed
        )
        file_hash = sha256_digest(command.content)
        plan_hash = _plan_hash(
            file_hash=file_hash,
            scope=scope,
            rows=classified,
            replace_rows=command.replace_row_numbers,
        )
        summary = _preview_summary(classified)
        return ParticipantSubmissionImportPreview(
            session_id=scope.session.session_id,
            scenario_snapshot_id=scope.snapshot.scenario_snapshot_id,
            configuration_version_id=scope.configuration.configuration_version_id,
            filename=filename,
            file_hash=file_hash,
            plan_hash=plan_hash,
            rows=classified,
            summary=summary,
            error_report_csv=_error_report(classified),
        )


def _load_scope(unit_of_work: UnitOfWork, session_id: str) -> _ImportScope:
    session = unit_of_work.session.get(session_id)
    if session is None:
        raise ParticipantSubmissionImportError("The selected session does not exist.")
    configuration = session.active_configuration
    if configuration is None or not configuration.is_activated:
        raise ParticipantSubmissionImportError(
            "The selected session has no activated configuration."
        )
    if configuration.response_format not in {
        ResponseFormat.PAIRWISE,
        ResponseFormat.DIRECT_RATING,
    }:
        raise ParticipantSubmissionImportError(
            "This response format is not supported for import. Direct ranking "
            "must be entered through a supported questionnaire workflow."
        )
    snapshot = unit_of_work.scenarios.get_by_id(configuration.scenario_snapshot_id)
    if snapshot is None:
        raise ParticipantSubmissionImportError(
            "The immutable scenario snapshot is unavailable."
        )
    scale = next(
        (item for item in snapshot.scales if item.scale_id == configuration.scale_id),
        None,
    )
    if scale is None or not scale.values:
        raise ParticipantSubmissionImportError(
            "The configured response scale is unavailable."
        )
    inputs: dict[str, list[str]] = {}
    for value in sorted(scale.values, key=lambda item: item.ordinal):
        inputs.setdefault(value.stable_value_key, []).append(value.scale_value_id)
        inputs.setdefault(value.label, []).append(value.scale_value_id)
    groups: dict[str, list[SessionStakeholderGroup]] = {}
    for group in configuration.stakeholder_groups:
        if not group.is_active:
            continue
        groups.setdefault(group.group_key, []).append(group)
        groups.setdefault(group.name, []).append(group)
    question_headers = {
        _question_header(question, snapshot): question
        for question in sorted(
            configuration.question_definitions,
            key=lambda item: (item.display_order, item.question_key),
        )
    }
    if len(question_headers) != len(configuration.question_definitions):
        raise ParticipantSubmissionImportError(
            "The frozen configuration produces duplicate import question headers."
        )
    return _ImportScope(
        session=session,
        configuration=configuration,
        snapshot=snapshot,
        scale_values_by_input={key: tuple(value) for key, value in inputs.items()},
        scale_choices=tuple(
            value.stable_value_key
            for value in sorted(scale.values, key=lambda item: item.ordinal)
        ),
        question_headers=question_headers,
        groups_by_input={key: tuple(value) for key, value in groups.items()},
    )


def _question_header(
    question: ResponseQuestionDefinition, snapshot: ScenarioSnapshot
) -> str:
    criteria = {item.criterion_id: item.criterion_key for item in snapshot.criteria}
    alternatives = {
        item.alternative_id: item.alternative_key for item in snapshot.alternatives
    }
    if question.question_type == QuestionType.CRITERION_PAIR:
        try:
            machine = (
                f"criterion_pair.{criteria[question.left_criterion_id or '']}."
                f"{criteria[question.right_criterion_id or '']}"
            )
        except KeyError as error:
            raise ParticipantSubmissionImportError(
                "A frozen pairwise question references an unknown criterion."
            ) from error
    elif question.question_type == QuestionType.CRITERION_RATING:
        try:
            machine = f"criterion_rating.{criteria[question.criterion_id or '']}"
        except KeyError as error:
            raise ParticipantSubmissionImportError(
                "A frozen rating question references an unknown criterion."
            ) from error
    elif question.question_type == QuestionType.ALTERNATIVE_RATING:
        try:
            machine = f"alternative_rating.{alternatives[question.alternative_id or '']}"
        except KeyError as error:
            raise ParticipantSubmissionImportError(
                "A frozen rating question references an unknown alternative."
            ) from error
    else:
        raise ParticipantSubmissionImportError(
            "This configuration contains a question type unsupported by import."
        )
    prompt = " ".join(question.prompt_snapshot.split())
    return f"answer.{machine} | {prompt}"


def _csv_template(headers: Sequence[str], example: Mapping[str, str] | None) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=headers, lineterminator="\n")
    writer.writeheader()
    if example is not None:
        writer.writerow({key: _csv_safe(example.get(key, "")) for key in headers})
    return stream.getvalue().encode("utf-8")


def _xlsx_template(
    scope: _ImportScope,
    headers: Sequence[str],
    example: Mapping[str, str] | None,
) -> bytes:
    workbook = Workbook()
    responses = workbook.active
    responses.title = "Responses"
    responses.append(list(headers))
    if example is not None:
        responses.append([example.get(header, "") for header in headers])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in responses[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    responses.freeze_panes = "A2"
    responses.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for index, header in enumerate(headers, start=1):
        responses.column_dimensions[get_column_letter(index)].width = min(
            55, max(14, len(header) * 0.9)
        )

    instructions = workbook.create_sheet("Instructions")
    instructions_rows = (
        ("Participant and submission import",),
        ("One row represents one participant and at most one submission.",),
        ("participant_ref is sensitive; it is never persisted in plaintext.",),
        ("record_state may be enrolled, draft, or submitted (blank means submitted).",),
        ("recorded_at must be a timezone-aware ISO 8601 timestamp and cannot be future.",),
        ("For pairwise questions, the selected value describes the frozen left criterion relative to the right criterion. Do not reverse a pair.",),
        ("Use stable scale keys or exact, unambiguous configured labels.",),
        ("Do not enter formulas. Formula cells are rejected.",),
    )
    for row in instructions_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 110
    for cell in instructions["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    instructions["A1"].font = Font(bold=True, size=14)

    scale_sheet = workbook.create_sheet("Scale values")
    scale_sheet.append(("Stable key", "Label", "Display value"))
    configured_scale = next(
        item for item in scope.snapshot.scales
        if item.scale_id == scope.configuration.scale_id
    )
    for value in sorted(configured_scale.values, key=lambda item: item.ordinal):
        scale_sheet.append(
            (value.stable_value_key, value.label, str(value.numeric_value or ""))
        )
    groups_sheet = workbook.create_sheet("Groups")
    groups_sheet.append(("Group key", "Label"))
    active_groups = sorted(
        (group for group in scope.configuration.stakeholder_groups if group.is_active),
        key=lambda item: (item.display_order, item.group_key),
    )
    for group in active_groups:
        groups_sheet.append((group.group_key, group.name))
    config_sheet = workbook.create_sheet("Configuration")
    for key, config_value in (
        ("template_schema_version", TEMPLATE_SCHEMA_VERSION),
        ("session_slug", scope.session.public_slug),
        ("session_title", scope.session.title),
        ("scenario_snapshot_version", scope.snapshot.declared_version),
        ("scenario_snapshot_hash", scope.snapshot.root_hash),
        ("configuration_version", str(scope.configuration.version_number)),
        ("configuration_hash", scope.configuration.config_hash),
        ("response_format", scope.configuration.response_format.value),
        ("response_target_type", scope.configuration.response_target_type.value),
    ):
        config_sheet.append((key, config_value))
    config_sheet.column_dimensions["A"].width = 32
    config_sheet.column_dimensions["B"].width = 72

    max_data_row = 1_001
    _add_dropdown(
        responses,
        headers,
        "group",
        f"'Groups'!$A$2:$A${max(2, len(active_groups) + 1)}",
        max_data_row,
    )
    _add_dropdown(responses, headers, "record_state", '"enrolled,draft,submitted"', max_data_row)
    for header in scope.question_headers:
        _add_dropdown(
            responses,
            headers,
            header,
            f"'Scale values'!$A$2:$A${len(configured_scale.values) + 1}",
            max_data_row,
        )
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _add_dropdown(
    worksheet: Any,
    headers: Sequence[str],
    header: str,
    formula: str,
    max_row: int,
) -> None:
    column = headers.index(header) + 1
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Choose a configured value from the list."
    validation.errorTitle = "Invalid import value"
    validation.prompt = "Use a configured stable key."
    validation.promptTitle = "Import value"
    worksheet.add_data_validation(validation)
    validation.add(
        f"{get_column_letter(column)}2:{get_column_letter(column)}{max_row}"
    )


def _example_row(scope: _ImportScope) -> dict[str, str]:
    active_groups = sorted(
        (group for group in scope.configuration.stakeholder_groups if group.is_active),
        key=lambda item: (item.display_order, item.group_key),
    )
    if not active_groups:
        raise ParticipantSubmissionImportError("The configuration has no active group.")
    row = {
        "template_schema_version": TEMPLATE_SCHEMA_VERSION,
        "session_slug": scope.session.public_slug,
        "configuration_version": str(scope.configuration.version_number),
        "participant_ref": "example-participant-001",
        "participant_alias": "Example participant 001",
        "participant_name": "",
        "participant_email": "",
        "group": active_groups[0].group_key,
        "record_state": ImportRecordState.SUBMITTED.value,
        "recorded_at": "",
    }
    for header in scope.question_headers:
        row[header] = scope.scale_choices[0]
    return row


def _validate_upload(filename: str, content: bytes, *, max_bytes: int) -> None:
    extension = PurePath(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ParticipantSubmissionImportError(
            "Upload a UTF-8 .csv file or an .xlsx workbook."
        )
    if not content:
        raise ParticipantSubmissionImportError("The uploaded file is empty.")
    if len(content) > max_bytes:
        raise ParticipantSubmissionImportError("The uploaded file exceeds the size limit.")
    if extension == ".xlsx" and not content.startswith(b"PK\x03\x04"):
        raise ParticipantSubmissionImportError("The .xlsx file signature is invalid.")
    if extension == ".csv" and (content.startswith(b"PK\x03\x04") or b"\x00" in content):
        raise ParticipantSubmissionImportError("The CSV content does not match its file type.")


def _read_upload(
    filename: str,
    content: bytes,
    *,
    scope: _ImportScope,
    max_rows: int,
) -> tuple[tuple[int, dict[str, object]], ...]:
    if PurePath(filename).suffix.lower() == ".csv":
        rows = _read_csv(content)
    else:
        rows = _read_xlsx(content)
    if len(rows) > max_rows:
        raise ParticipantSubmissionImportError("The uploaded file exceeds the row limit.")
    if not rows:
        raise ParticipantSubmissionImportError("The uploaded file contains no data rows.")
    expected = {*RESERVED_COLUMNS, *scope.question_headers}
    actual = set(rows[0][1])
    missing = sorted(expected - actual)
    extra = sorted(actual - expected)
    if missing or extra:
        details = []
        if missing:
            details.append(f"missing {len(missing)} required column(s)")
        if extra:
            details.append(f"contains {len(extra)} unknown column(s)")
        raise ParticipantSubmissionImportError(
            "The response columns do not match the frozen template: "
            + " and ".join(details)
            + "."
        )
    return rows


def _read_csv(content: bytes) -> tuple[tuple[int, dict[str, object]], ...]:
    try:
        text = content.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as error:
        raise ParticipantSubmissionImportError("CSV files must use UTF-8 encoding.") from error
    try:
        reader = csv.DictReader(io.StringIO(text, newline=""), strict=True)
        if reader.fieldnames is None or any(not name for name in reader.fieldnames):
            raise ParticipantSubmissionImportError("The CSV header is invalid.")
        if len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise ParticipantSubmissionImportError("CSV column names must be unique.")
        return tuple(
            (number, {key: value for key, value in row.items() if key is not None})
            for number, row in enumerate(reader, start=2)
            if any((value or "").strip() for value in row.values() if value is not None)
        )
    except csv.Error as error:
        raise ParticipantSubmissionImportError("The CSV file is malformed.") from error


def _read_xlsx(content: bytes) -> tuple[tuple[int, dict[str, object]], ...]:
    _inspect_xlsx_archive(content)
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except Exception as error:
        raise ParticipantSubmissionImportError("The .xlsx workbook is malformed.") from error
    try:
        if "Responses" not in workbook.sheetnames:
            raise ParticipantSubmissionImportError(
                "The workbook must contain a Responses sheet."
            )
        sheet = workbook["Responses"]
        iterator = sheet.iter_rows(values_only=False)
        try:
            header_cells = next(iterator)
        except StopIteration as error:
            raise ParticipantSubmissionImportError("The Responses sheet is empty.") from error
        headers = tuple("" if cell.value is None else str(cell.value).strip() for cell in header_cells)
        if any(not header for header in headers) or len(headers) != len(set(headers)):
            raise ParticipantSubmissionImportError(
                "Responses sheet column names must be nonempty and unique."
            )
        rows: list[tuple[int, dict[str, object]]] = []
        for row_number, cells in enumerate(iterator, start=2):
            values: list[object] = []
            for cell in cells[: len(headers)]:
                if cell.data_type == "f" or (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                ):
                    raise ParticipantSubmissionImportError(
                        f"Formula cells are not permitted (Responses row {row_number})."
                    )
                values.append(cell.value)
            values.extend([None] * (len(headers) - len(values)))
            if any(value is not None and str(value).strip() for value in values):
                rows.append((row_number, dict(zip(headers, values, strict=True))))
        return tuple(rows)
    finally:
        workbook.close()


def _inspect_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            total = 0
            names = set()
            for item in archive.infolist():
                normalized = PurePath(item.filename.replace("\\", "/"))
                if normalized.is_absolute() or ".." in normalized.parts:
                    raise ParticipantSubmissionImportError(
                        "The workbook contains an unsafe archive path."
                    )
                lowered = item.filename.casefold()
                if lowered in names:
                    raise ParticipantSubmissionImportError(
                        "The workbook contains duplicate archive members."
                    )
                names.add(lowered)
                if (
                    "vbaproject" in lowered
                    or lowered.startswith("xl/externallinks/")
                    or lowered.endswith((".bin", ".vba"))
                ):
                    raise ParticipantSubmissionImportError(
                        "Macros and external workbook links are not permitted."
                    )
                if item.file_size > _MAX_XLSX_MEMBER_BYTES:
                    raise ParticipantSubmissionImportError(
                        "The workbook contains an oversized component."
                    )
                total += item.file_size
                if total > _MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise ParticipantSubmissionImportError(
                        "The workbook expands beyond the safe processing limit."
                    )
                if item.compress_size and item.file_size / item.compress_size > 200:
                    raise ParticipantSubmissionImportError(
                        "The workbook contains an unsafe compression ratio."
                    )
                if lowered == "xl/calcchain.xml":
                    raise ParticipantSubmissionImportError(
                        "Workbooks requiring formula evaluation are not permitted."
                    )
                if lowered.endswith((".xml", ".rels")) and item.file_size:
                    payload = archive.read(item)
                    if (
                        lowered.endswith(".rels")
                        and b'TargetMode="External"' in payload
                    ):
                        raise ParticipantSubmissionImportError(
                            "External workbook links are not permitted."
                        )
                    if (
                        lowered.startswith("xl/worksheets/")
                        and re.search(br"<f(?:\s|>)", payload) is not None
                    ):
                        raise ParticipantSubmissionImportError(
                            "Formula cells are not permitted."
                        )
    except zipfile.BadZipFile as error:
        raise ParticipantSubmissionImportError("The .xlsx archive is malformed.") from error


def _parse_source_row(
    row_number: int,
    source: Mapping[str, object],
    *,
    scope: _ImportScope,
    protector: IdentityProtector,
    now: datetime,
) -> _ParsedRow:
    issues: list[ImportIssue] = []
    schema = _cell_text(source.get("template_schema_version"))
    slug = _cell_text(source.get("session_slug"))
    version = _cell_text(source.get("configuration_version"))
    _expect_binding(issues, row_number, "template_schema_version", schema, TEMPLATE_SCHEMA_VERSION)
    _expect_binding(issues, row_number, "session_slug", slug, scope.session.public_slug)
    _expect_binding(
        issues, row_number, "configuration_version", version,
        str(scope.configuration.version_number),
    )
    raw_ref = _cell_text(source.get("participant_ref"))
    normalized_ref = normalize_participant_reference(raw_ref)
    reference_digest = None
    if not normalized_ref:
        issues.append(_issue(row_number, "participant_ref", "reference.required", "participant_ref is required."))
    elif len(normalized_ref) > 320:
        issues.append(_issue(row_number, "participant_ref", "reference.too_long", "participant_ref is too long."))
    else:
        reference_digest = protector.participant_reference_digest(normalized_ref)
    alias = _cell_text(source.get("participant_alias"))
    if not alias:
        issues.append(_issue(row_number, "participant_alias", "alias.required", "participant_alias is required."))
    elif len(alias) > 160:
        issues.append(_issue(row_number, "participant_alias", "alias.too_long", "participant_alias must be 160 characters or fewer."))
    name = _optional_cell_text(source.get("participant_name"))
    email = _optional_cell_text(source.get("participant_email"))
    if (name or email) and scope.session.identity_policy.casefold() == "anonymous":
        issues.append(_issue(row_number, "participant_name", "identity.anonymous_session", "Identity data cannot be imported into an anonymous session."))
    group_input = _cell_text(source.get("group"))
    groups = scope.groups_by_input.get(group_input, ())
    group = groups[0] if len(groups) == 1 else None
    if not group_input:
        active = tuple({item.session_stakeholder_group_id: item for values in scope.groups_by_input.values() for item in values}.values())
        if len(active) == 1:
            group = active[0]
        else:
            issues.append(_issue(row_number, "group", "group.required", "group is required.", tuple(sorted({item.group_key for values in scope.groups_by_input.values() for item in values}))))
    elif not groups:
        issues.append(_issue(row_number, "group", "group.unknown", "group is not active in the frozen configuration.", tuple(sorted({item.group_key for values in scope.groups_by_input.values() for item in values}))))
    elif len(groups) > 1:
        issues.append(_issue(row_number, "group", "group.ambiguous", "group matches more than one configured group; use its stable key."))
    state_input = _cell_text(source.get("record_state")) or ImportRecordState.SUBMITTED.value
    try:
        state = ImportRecordState(state_input.casefold())
    except ValueError:
        state = ImportRecordState.SUBMITTED
        issues.append(_issue(row_number, "record_state", "record_state.invalid", "record_state must be enrolled, draft, or submitted.", tuple(item.value for item in ImportRecordState)))
    recorded_at = _parse_recorded_at(source.get("recorded_at"), row_number, issues, now)
    answers: list[tuple[ResponseQuestionDefinition, str]] = []
    for header, question in scope.question_headers.items():
        entered = _cell_text(source.get(header))
        if not entered:
            continue
        matches = scope.scale_values_by_input.get(entered, ())
        if not matches:
            issues.append(_issue(row_number, header, "scale_value.unknown", "The response is not in the configured scale.", scope.scale_choices))
        elif len(matches) > 1:
            issues.append(_issue(row_number, header, "scale_value.ambiguous", "The scale label is ambiguous; use a stable scale-value key.", scope.scale_choices))
        else:
            answers.append((question, matches[0]))
    if state == ImportRecordState.ENROLLED and answers:
        issues.append(_issue(row_number, "record_state", "enrolled.answers_present", "Enrolled-only rows must leave every answer blank."))
    if state == ImportRecordState.SUBMITTED and not scope.configuration.allow_incomplete_submission:
        answered = {question.question_definition_id for question, _ in answers}
        missing = tuple(
            question for question in scope.configuration.question_definitions
            if question.required and question.question_definition_id not in answered
        )
        if missing:
            issues.append(_issue(row_number, "record_state", "answers.required_missing", f"Submitted row is missing {len(missing)} required answer(s)."))
    return _ParsedRow(
        row_number=row_number,
        participant_ref=raw_ref,
        reference_digest=reference_digest,
        participant_alias=alias,
        participant_name=name,
        participant_email=email,
        group=group,
        record_state=state,
        recorded_at=recorded_at,
        answers=tuple(answers),
        issues=tuple(issues),
    )


def _classify_row(
    row: _ParsedRow,
    *,
    scope: _ImportScope,
    duplicate: bool,
    existing_participant_id: str | None,
    participant_by_id: Mapping[str, Participant],
    replace_selected: bool,
) -> ParticipantSubmissionImportRow:
    issues = list(row.issues)
    if duplicate:
        issues.append(_issue(row.row_number, "participant_ref", "reference.duplicate_in_file", "participant_ref occurs more than once in this file."))
    existing = (
        None if existing_participant_id is None
        else participant_by_id.get(existing_participant_id)
    )
    unresolvable = False
    if existing is not None and (
        existing.session_id != scope.session.session_id
        or existing.configuration_version_id != scope.configuration.configuration_version_id
        or row.group is None
        or existing.session_stakeholder_group_id != row.group.session_stakeholder_group_id
    ):
        unresolvable = True
        issues.append(_issue(row.row_number, "participant_ref", "conflict.binding_mismatch", "The matched participant is bound to a different session configuration or stakeholder group."))
    blocking = any(issue.blocking for issue in issues)
    if blocking or unresolvable:
        status = ImportRowStatus.BLOCKING_ERROR
    elif existing_participant_id is None:
        status = ImportRowStatus.READY
    elif replace_selected:
        status = ImportRowStatus.PLANNED_REPLACEMENT
    else:
        status = ImportRowStatus.SKIPPED_CONFLICT
        issues.append(ImportIssue(row.row_number, "participant_ref", "conflict.skipped", "An existing participant matches this reference; the row will be skipped unless replacement is selected.", blocking=False))
    if row.recorded_at is None:
        issues.append(ImportIssue(row.row_number, "recorded_at", "recorded_at.defaulted", "recorded_at is blank; the actual import time will be used.", blocking=False))
    return ParticipantSubmissionImportRow(
        row_number=row.row_number,
        participant_ref=row.participant_ref,
        participant_alias=row.participant_alias,
        group_key=None if row.group is None else row.group.group_key,
        group_name=None if row.group is None else row.group.name,
        record_state=row.record_state,
        recorded_at=row.recorded_at,
        answer_count=len(row.answers),
        has_identity=bool(row.participant_name or row.participant_email),
        status=status,
        existing_participant_id=existing_participant_id,
        issues=tuple(issues),
    )


def _preview_summary(rows: Sequence[ParticipantSubmissionImportRow]) -> ImportPreviewSummary:
    applied = tuple(row for row in rows if row.status in {ImportRowStatus.READY, ImportRowStatus.PLANNED_REPLACEMENT})
    return ImportPreviewSummary(
        total_rows=len(rows),
        new_participants=sum(row.status == ImportRowStatus.READY for row in rows),
        enrolled_only=sum(row.record_state == ImportRecordState.ENROLLED for row in applied),
        drafts=sum(row.record_state == ImportRecordState.DRAFT for row in applied),
        submitted=sum(row.record_state == ImportRecordState.SUBMITTED for row in applied),
        skipped_conflicts=sum(row.status == ImportRowStatus.SKIPPED_CONFLICT for row in rows),
        planned_replacements=sum(row.status == ImportRowStatus.PLANNED_REPLACEMENT for row in rows),
        identity_rows=sum(row.has_identity for row in applied),
        warning_count=sum(not issue.blocking for row in rows for issue in row.issues),
        blocking_error_count=sum(issue.blocking for row in rows for issue in row.issues),
    )


def _plan_hash(
    *,
    file_hash: str,
    scope: _ImportScope,
    rows: Sequence[ParticipantSubmissionImportRow],
    replace_rows: frozenset[int],
) -> str:
    return hash_json({
        "schema": TEMPLATE_SCHEMA_VERSION,
        "file_hash": file_hash,
        "session_id": scope.session.session_id,
        "scenario_snapshot_id": scope.snapshot.scenario_snapshot_id,
        "configuration_version_id": scope.configuration.configuration_version_id,
        "configuration_hash": scope.configuration.config_hash,
        "replace_rows": sorted(replace_rows),
        "rows": [
            {"row": row.row_number, "status": row.status.value, "state": row.record_state.value, "answers": row.answer_count}
            for row in rows
        ],
    })


def _error_report(rows: Sequence[ParticipantSubmissionImportRow]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerow(("row_number", "participant_ref", "column", "issue_code", "message", "accepted_values"))
    for row in rows:
        for issue in row.issues:
            writer.writerow((
                row.row_number,
                _csv_safe(row.participant_ref),
                issue.column,
                issue.code,
                issue.message,
                _csv_safe("; ".join(issue.accepted_values)),
            ))
    return stream.getvalue().encode("utf-8")


def _expect_binding(
    issues: list[ImportIssue], row: int, column: str, actual: str, expected: str
) -> None:
    if actual != expected:
        issues.append(_issue(row, column, f"binding.{column}", "The file is bound to a different session or configuration."))


def _parse_recorded_at(
    value: object,
    row: int,
    issues: list[ImportIssue],
    now: datetime,
) -> datetime | None:
    text = _cell_text(value)
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        parsed = None
    if parsed is None or parsed.tzinfo is None or parsed.utcoffset() is None:
        issues.append(_issue(row, "recorded_at", "recorded_at.invalid", "recorded_at must be a timezone-aware ISO 8601 timestamp."))
        return None
    parsed = parsed.astimezone(UTC)
    if parsed > now:
        issues.append(_issue(row, "recorded_at", "recorded_at.future", "recorded_at cannot be in the future."))
    return parsed


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def _optional_cell_text(value: object) -> str | None:
    result = _cell_text(value)
    return result or None


def _issue(
    row: int, column: str, code: str, message: str,
    accepted: tuple[str, ...] = (),
) -> ImportIssue:
    return ImportIssue(row, column, code, message, accepted)


def _csv_safe(value: str) -> str:
    return "'" + value if value.startswith(_FORMULA_PREFIXES) else value


def _safe_filename(value: str) -> str:
    result = _SAFE_FILENAME.sub("-", value).strip(".-")[:220]
    return result or "participant-import"


def _safe_uploaded_filename(value: str) -> str:
    name = PurePath(value).name
    stem = _safe_filename(PurePath(name).stem)
    suffix = PurePath(name).suffix.lower()
    return f"{stem}{suffix}"


@dataclass(frozen=True, slots=True, repr=False)
class ApplyParticipantSubmissionImportCommand:
    session_id: str
    filename: str
    content: bytes = field(repr=False)
    expected_file_hash: str
    expected_plan_hash: str
    actor_id: str
    actor_roles: frozenset[str]
    replace_row_numbers: frozenset[int] = frozenset()
    lifecycle_override: bool = False
    lifecycle_override_reason: str | None = None
    resubmission_override: bool = False
    resubmission_override_reason: str | None = None
    identity_processing_attested: bool = False
    identity_processing_basis: str | None = None
    actor_type: ActorType = ActorType.USER
    correlation_id: str = field(default_factory=lambda: str(new_id()))

    def __post_init__(self) -> None:
        for value, label in (
            (self.session_id, "session_id"),
            (self.filename, "filename"),
            (self.expected_file_hash, "expected_file_hash"),
            (self.expected_plan_hash, "expected_plan_hash"),
            (self.actor_id, "actor_id"),
            (self.correlation_id, "correlation_id"),
        ):
            if not value.strip():
                raise ParticipantSubmissionImportError(f"{label} cannot be empty.")
        if self.actor_type != ActorType.USER or "admin" not in self.actor_roles:
            raise ParticipantSubmissionImportError(
                "Participant imports require an authenticated authorized administrator."
            )

    def __repr__(self) -> str:
        return (
            "ApplyParticipantSubmissionImportCommand("
            f"session_id={self.session_id!r}, filename={self.filename!r}, "
            "content=<redacted>, identity=<redacted>)"
        )


@dataclass(frozen=True, slots=True, repr=False)
class AppliedParticipantImportRow:
    row_number: int
    participant_ref: str
    participant_id: str | None
    submission_id: str | None
    status: str
    message: str


@dataclass(frozen=True, slots=True, repr=False)
class ApplyParticipantSubmissionImportResult:
    batch_id: str
    file_hash: str
    plan_hash: str
    created_participant_count: int
    enrolled_count: int
    draft_count: int
    submitted_count: int
    skipped_count: int
    replaced_count: int
    identity_count: int
    imported_at: datetime
    rows: tuple[AppliedParticipantImportRow, ...]
    result_report_csv: bytes = field(repr=False)
    already_applied: bool = False

    def __repr__(self) -> str:
        return (
            "ApplyParticipantSubmissionImportResult("
            f"batch_id={self.batch_id!r}, row_count={len(self.rows)}, "
            "sensitive_values=<redacted>)"
        )


class ApplyParticipantSubmissionImport:
    """Reverify and atomically apply a previewed import plan."""

    def __init__(
        self,
        unit_of_work_factory: UnitOfWorkFactory,
        identity_protector: IdentityProtector,
        *,
        max_bytes: int,
        max_rows: int,
        identity_retention_days: int,
        clock: Clock = utc_now,
        id_factory: IdFactory = lambda: str(new_id()),
    ) -> None:
        self._unit_of_work_factory = unit_of_work_factory
        self._identity_protector = identity_protector
        self._max_rows = max_rows
        self._identity_retention_days = identity_retention_days
        self._clock = clock
        self._id_factory = id_factory
        self._preview = PreviewParticipantSubmissionImport(
            unit_of_work_factory,
            identity_protector,
            max_bytes=max_bytes,
            max_rows=max_rows,
            clock=clock,
        )

    def execute(
        self, command: ApplyParticipantSubmissionImportCommand
    ) -> ApplyParticipantSubmissionImportResult:
        actual_file_hash = sha256_digest(command.content)
        if actual_file_hash != command.expected_file_hash:
            raise ParticipantSubmissionImportError(
                "The uploaded content changed after preview. Preview it again."
            )
        # A committed batch is the retry authority. Conflict classification
        # necessarily changes after its participants exist, so idempotent
        # retries must be recognized before rebuilding the preview plan.
        with self._unit_of_work_factory() as unit_of_work:
            applied_batch = unit_of_work.participant_imports.get_batch(
                command.session_id,
                command.expected_file_hash,
                command.expected_plan_hash,
            )
        if applied_batch is not None:
            return _already_applied_result(applied_batch)
        preview = self._preview.execute(
            PreviewParticipantSubmissionImportCommand(
                session_id=command.session_id,
                filename=command.filename,
                content=command.content,
                replace_row_numbers=command.replace_row_numbers,
            )
        )
        if preview.file_hash != command.expected_file_hash:
            raise ParticipantSubmissionImportError(
                "The uploaded content changed after preview. Preview it again."
            )
        if preview.plan_hash != command.expected_plan_hash:
            raise ParticipantSubmissionImportError(
                "The selected import plan changed after preview. Preview it again."
            )
        if preview.has_blocking_errors:
            raise ParticipantSubmissionImportError(
                "Blocking preview errors must be resolved before apply."
            )
        imported_at = self._clock()
        filename = _safe_uploaded_filename(command.filename)
        with self._unit_of_work_factory() as unit_of_work:
            locked_session = unit_of_work.session.get_for_update(command.session_id)
            if locked_session is None:
                raise ParticipantSubmissionImportError(
                    "The selected session no longer exists."
                )
            _require_apply_lifecycle(locked_session, command)
            scope = _load_scope(unit_of_work, command.session_id)
            if (
                scope.configuration.configuration_version_id
                != preview.configuration_version_id
                or scope.snapshot.scenario_snapshot_id
                != preview.scenario_snapshot_id
            ):
                raise ParticipantSubmissionImportError(
                    "The active frozen configuration changed after preview."
                )
            existing_batch = unit_of_work.participant_imports.get_batch(
                command.session_id, preview.file_hash, preview.plan_hash
            )
            if existing_batch is not None:
                return _already_applied_result(existing_batch)

            source_rows = _read_upload(
                filename,
                command.content,
                scope=scope,
                max_rows=self._max_rows,
            )
            parsed_by_row = {
                row.row_number: row
                for row in (
                    _parse_source_row(
                        row_number,
                        source,
                        scope=scope,
                        protector=self._identity_protector,
                        now=imported_at,
                    )
                    for row_number, source in source_rows
                )
            }
            if preview.summary.identity_rows:
                if not command.identity_processing_attested:
                    raise ParticipantSubmissionImportError(
                        "Identity-processing authorization must be attested before apply."
                    )
                if not (command.identity_processing_basis or "").strip():
                    raise ParticipantSubmissionImportError(
                        "Identity-processing attestation requires a processing basis."
                    )
            if preview.summary.planned_replacements and not scope.configuration.allow_resubmissions:
                if not command.resubmission_override:
                    raise ParticipantSubmissionImportError(
                        "Replacement requires a separate resubmission-policy override."
                    )
                if not (command.resubmission_override_reason or "").strip():
                    raise ParticipantSubmissionImportError(
                        "Resubmission-policy override requires a reason."
                    )

            current_mapping = unit_of_work.participant_imports.participants_by_reference_digests(
                command.session_id,
                tuple(
                    row.reference_digest
                    for row in parsed_by_row.values()
                    if row.reference_digest is not None
                ),
            )
            _verify_concurrent_conflicts(preview, parsed_by_row, current_mapping)
            batch_id = self._id_factory()
            batch = ParticipantSubmissionImportBatch(
                batch_id=batch_id,
                session_id=command.session_id,
                scenario_snapshot_id=scope.snapshot.scenario_snapshot_id,
                configuration_version_id=scope.configuration.configuration_version_id,
                file_hash=preview.file_hash,
                plan_hash=preview.plan_hash,
                filename=filename,
                row_count=preview.summary.total_rows,
                created_participant_count=preview.summary.new_participants,
                enrolled_count=preview.summary.enrolled_only,
                draft_count=preview.summary.drafts,
                submitted_count=preview.summary.submitted,
                skipped_count=preview.summary.skipped_conflicts,
                replaced_count=preview.summary.planned_replacements,
                identity_count=preview.summary.identity_rows,
                lifecycle_override=command.lifecycle_override,
                resubmission_override=command.resubmission_override,
                identity_attested=command.identity_processing_attested,
                terminal_status="applied",
                imported_at=imported_at,
                imported_by=command.actor_id,
                correlation_id=command.correlation_id,
            )
            unit_of_work.participant_imports.add_batch(batch)
            result_rows: list[AppliedParticipantImportRow] = []
            for planned in preview.rows:
                parsed = parsed_by_row[planned.row_number]
                if planned.status == ImportRowStatus.SKIPPED_CONFLICT:
                    result_rows.append(AppliedParticipantImportRow(
                        row_number=planned.row_number,
                        participant_ref=parsed.participant_ref,
                        participant_id=planned.existing_participant_id,
                        submission_id=None,
                        status="skipped_conflict",
                        message="Existing participant preserved; no changes applied.",
                    ))
                    continue
                if planned.status == ImportRowStatus.BLOCKING_ERROR:
                    raise ParticipantSubmissionImportError(
                        "The import plan contains a blocking row."
                    )
                participant, reference, created = self._participant_for_row(
                    unit_of_work,
                    scope=scope,
                    parsed=parsed,
                    planned=planned,
                    batch_id=batch_id,
                    actor_id=command.actor_id,
                    imported_at=imported_at,
                )
                submission, predecessor_id = self._submission_for_row(
                    unit_of_work,
                    scope=scope,
                    participant=participant,
                    parsed=parsed,
                    replacing=planned.status == ImportRowStatus.PLANNED_REPLACEMENT,
                    resubmission_override=command.resubmission_override,
                    actor_id=command.actor_id,
                    imported_at=imported_at,
                )
                if parsed.participant_name or parsed.participant_email:
                    self._store_identity(
                        unit_of_work,
                        participant=participant,
                        parsed=parsed,
                        batch_id=batch_id,
                        command=command,
                        imported_at=imported_at,
                    )
                import_record_id = self._id_factory()
                unit_of_work.participant_imports.add_record(
                    ParticipantSubmissionImportRecord(
                        import_record_id=import_record_id,
                        batch_id=batch_id,
                        reference_mapping_id=reference.reference_mapping_id,
                        participant_id=participant.participant_id,
                        submission_id=None if submission is None else submission.submission_id,
                        predecessor_submission_id=predecessor_id,
                        source_row_number=parsed.row_number,
                        record_state=parsed.record_state.value,
                        recorded_at=parsed.recorded_at or imported_at,
                        imported_at=imported_at,
                    )
                )
                unit_of_work.participant_imports.add_consent_disposition(
                    AdminImportConsentDisposition(
                        disposition_id=self._id_factory(),
                        participant_id=participant.participant_id,
                        submission_id=None if submission is None else submission.submission_id,
                        configuration_version_id=scope.configuration.configuration_version_id,
                        batch_id=batch_id,
                        administrator_id=command.actor_id,
                        reason_code="not_applicable_admin_import",
                        recorded_at=imported_at,
                    )
                )
                unit_of_work.audit_events.add(
                    operational_audit_event(
                        event_id=self._id_factory(),
                        occurred_at=imported_at,
                        session_id=command.session_id,
                        actor_id=command.actor_id,
                        actor_type=command.actor_type,
                        action=(
                            AuditAction.UPDATED
                            if planned.status
                            == ImportRowStatus.PLANNED_REPLACEMENT
                            else AuditAction.CREATED
                        ),
                        entity_type="participant_submission_import_record",
                        entity_id=import_record_id,
                        correlation_id=command.correlation_id,
                        use_case="apply_participant_submission_import",
                        reason_text=(
                            command.resubmission_override_reason
                            if planned.status
                            == ImportRowStatus.PLANNED_REPLACEMENT
                            else None
                        ),
                        after_json={
                            "participant_id": participant.participant_id,
                            "submission_id": (
                                None
                                if submission is None
                                else submission.submission_id
                            ),
                            "predecessor_submission_id": predecessor_id,
                            "record_state": parsed.record_state.value,
                            "recorded_at": parsed.recorded_at or imported_at,
                            "replacement": planned.status
                            == ImportRowStatus.PLANNED_REPLACEMENT,
                            "consent_disposition": (
                                "not_applicable_admin_import"
                            ),
                        },
                    )
                )
                result_rows.append(AppliedParticipantImportRow(
                    row_number=parsed.row_number,
                    participant_ref=parsed.participant_ref,
                    participant_id=participant.participant_id,
                    submission_id=None if submission is None else submission.submission_id,
                    status=("replaced" if not created else parsed.record_state.value),
                    message="Consent not applicable — administrator import",
                ))

            unit_of_work.audit_events.add(
                operational_audit_event(
                    event_id=self._id_factory(),
                    occurred_at=imported_at,
                    session_id=command.session_id,
                    actor_id=command.actor_id,
                    actor_type=command.actor_type,
                    action=AuditAction.CREATED,
                    entity_type="participant_submission_import_batch",
                    entity_id=batch_id,
                    correlation_id=command.correlation_id,
                    use_case="apply_participant_submission_import",
                    reason_text=(
                        command.lifecycle_override_reason
                        if command.lifecycle_override
                        else None
                    ),
                    after_json={
                        "row_count": preview.summary.total_rows,
                        "created_participant_count": preview.summary.new_participants,
                        "enrolled_count": preview.summary.enrolled_only,
                        "draft_count": preview.summary.drafts,
                        "submitted_count": preview.summary.submitted,
                        "skipped_count": preview.summary.skipped_conflicts,
                        "replaced_count": preview.summary.planned_replacements,
                        "identity_count": preview.summary.identity_rows,
                        "lifecycle_override": command.lifecycle_override,
                        "resubmission_override": command.resubmission_override,
                        "identity_attested": command.identity_processing_attested,
                    },
                )
            )
            unit_of_work.commit()

        rows = tuple(result_rows)
        return ApplyParticipantSubmissionImportResult(
            batch_id=batch_id,
            file_hash=preview.file_hash,
            plan_hash=preview.plan_hash,
            created_participant_count=preview.summary.new_participants,
            enrolled_count=preview.summary.enrolled_only,
            draft_count=preview.summary.drafts,
            submitted_count=preview.summary.submitted,
            skipped_count=preview.summary.skipped_conflicts,
            replaced_count=preview.summary.planned_replacements,
            identity_count=preview.summary.identity_rows,
            imported_at=imported_at,
            rows=rows,
            result_report_csv=_result_report(rows),
        )

    def _participant_for_row(
        self,
        unit_of_work: UnitOfWork,
        *,
        scope: _ImportScope,
        parsed: _ParsedRow,
        planned: ParticipantSubmissionImportRow,
        batch_id: str,
        actor_id: str,
        imported_at: datetime,
    ) -> tuple[Participant, ParticipantImportReference, bool]:
        if parsed.group is None or parsed.reference_digest is None:
            raise ParticipantSubmissionImportError("A ready import row is incomplete.")
        if planned.existing_participant_id is not None:
            participant = unit_of_work.participants.get_for_update(
                planned.existing_participant_id
            )
            reference = unit_of_work.participant_imports.get_reference(
                scope.session.session_id, parsed.reference_digest
            )
            if participant is None or reference is None:
                raise ParticipantSubmissionImportError(
                    "A participant conflict changed during apply. Preview again."
                )
            participant.validate_binding(
                configuration=scope.configuration, group=parsed.group
            )
            return participant, reference, False
        recorded_at = parsed.recorded_at or imported_at
        participant = Participant.enroll(
            participant_id=self._id_factory(),
            session_id=scope.session.session_id,
            configuration=scope.configuration,
            group=parsed.group,
            actor_id=actor_id,
            at=recorded_at,
            alias=parsed.participant_alias,
        )
        unit_of_work.participants.add(participant)
        reference = ParticipantImportReference(
            reference_mapping_id=self._id_factory(),
            session_id=scope.session.session_id,
            participant_id=participant.participant_id,
            reference_digest=parsed.reference_digest,
            first_batch_id=batch_id,
            created_at=imported_at,
        )
        unit_of_work.participant_imports.add_reference(reference)
        return participant, reference, True

    def _submission_for_row(
        self,
        unit_of_work: UnitOfWork,
        *,
        scope: _ImportScope,
        participant: Participant,
        parsed: _ParsedRow,
        replacing: bool,
        resubmission_override: bool,
        actor_id: str,
        imported_at: datetime,
    ) -> tuple[Submission | None, str | None]:
        if parsed.record_state == ImportRecordState.ENROLLED:
            return None, None
        plan = unit_of_work.submissions.lock_attempt_plan(
            participant.participant_id,
            scope.configuration.configuration_version_id,
        )
        predecessor: Submission | None
        if plan.draft is not None:
            if not replacing:
                raise ParticipantSubmissionImportError(
                    "A draft conflict changed during apply. Preview again."
                )
            withdrawn = plan.draft.withdraw_draft_for_replacement(
                actor_id=actor_id,
                reason="Replaced by administrator import",
                at=imported_at,
            )
            unit_of_work.submissions.save(withdrawn)
            predecessor = withdrawn
        else:
            predecessor = plan.previous_submission
        recorded_at = parsed.recorded_at or imported_at
        attempt_at = recorded_at
        if predecessor is not None and attempt_at < predecessor.updated_at:
            attempt_at = imported_at
        if attempt_at < participant.updated_at:
            attempt_at = imported_at
        if participant.joined_at is None:
            participant = participant.join(actor_id=actor_id, at=attempt_at)
        if participant.started_at is None:
            participant = participant.start(actor_id=actor_id, at=attempt_at)
        unit_of_work.participants.save(participant)
        submission = Submission.start(
            submission_id=self._id_factory(),
            participant=participant,
            configuration=scope.configuration,
            group=parsed.group or _unreachable_group(),
            previous_submission=predecessor,
            administrative_resubmission_override=(
                replacing
                and (
                    resubmission_override
                    or scope.configuration.allow_resubmissions
                )
            ),
            actor_id=actor_id,
            at=attempt_at,
            client_metadata_json={
                "authorship": "moderator_captured_import",
                "recorded_at": recorded_at.isoformat(),
            },
        )
        for question, scale_value_id in parsed.answers:
            answer = SubmissionAnswer.create(
                submission_answer_id=self._id_factory(),
                submission_id=submission.submission_id,
                question=question,
                raw_value_json={"selected_scale_value_id": scale_value_id},
                answered_at=attempt_at,
            )
            submission = submission.save_answer(
                answer,
                configuration=scope.configuration,
                actor_id=actor_id,
                at=attempt_at,
            )
        if parsed.record_state == ImportRecordState.SUBMITTED:
            submission = submission.submit(
                scope.configuration, actor_id=actor_id, at=attempt_at
            )
            if predecessor is not None and predecessor.status == SubmissionStatus.SUBMITTED:
                superseded = predecessor.supersede_with(
                    submission, actor_id=actor_id, at=imported_at
                )
                unit_of_work.submissions.save(superseded)
            if participant.submitted_at is None:
                participant = participant.record_submission(
                    actor_id=actor_id, at=attempt_at
                ).complete(actor_id=actor_id, at=attempt_at)
                unit_of_work.participants.save(participant)
        unit_of_work.submissions.add(submission)
        return submission, None if predecessor is None else predecessor.submission_id

    def _store_identity(
        self,
        unit_of_work: UnitOfWork,
        *,
        participant: Participant,
        parsed: _ParsedRow,
        batch_id: str,
        command: ApplyParticipantSubmissionImportCommand,
        imported_at: datetime,
    ) -> None:
        if unit_of_work.participant_identities.get(participant.participant_id) is not None:
            raise ParticipantSubmissionImportError(
                "The matched participant already has protected identity data."
            )
        context = f"participant-identity:{participant.participant_id}"
        retention_until = imported_at + timedelta(
            days=self._identity_retention_days
        )
        identity = ParticipantIdentity(
            participant_id=participant.participant_id,
            consent_version="not_applicable_admin_import",
            consented_at=imported_at,
            retention_until=retention_until,
            display_name_ciphertext=(
                None
                if parsed.participant_name is None
                else self._identity_protector.encrypt(
                    parsed.participant_name, context=context + ":name"
                )
            ),
            email_ciphertext=(
                None
                if parsed.participant_email is None
                else self._identity_protector.encrypt(
                    parsed.participant_email, context=context + ":email"
                )
            ),
            email_lookup_hash=(
                None
                if parsed.participant_email is None
                else self._identity_protector.email_lookup_digest(
                    normalize_email_for_lookup(parsed.participant_email)
                )
            ),
        )
        unit_of_work.participant_identities.add(identity)
        unit_of_work.participant_imports.add_identity_authority(
            ImportedIdentityAuthority(
                authority_id=self._id_factory(),
                participant_id=participant.participant_id,
                batch_id=batch_id,
                administrator_id=command.actor_id,
                processing_basis=(command.identity_processing_basis or "").strip(),
                attested_at=imported_at,
                retention_until=retention_until,
            )
        )


def _require_apply_lifecycle(
    session: Session, command: ApplyParticipantSubmissionImportCommand
) -> None:
    if session.status in {SessionStatus.CANCELED, SessionStatus.ARCHIVED}:
        raise ParticipantSubmissionImportError(
            "Canceled and archived sessions cannot be changed by import."
        )
    if session.status in {SessionStatus.PAUSED, SessionStatus.CLOSED}:
        if not command.lifecycle_override:
            raise ParticipantSubmissionImportError(
                "Paused or closed sessions require a separate historical-data override."
            )
        if not (command.lifecycle_override_reason or "").strip():
            raise ParticipantSubmissionImportError(
                "The lifecycle override requires an audit reason."
            )
        return
    if session.status not in {SessionStatus.DRAFT, SessionStatus.OPEN}:
        raise ParticipantSubmissionImportError(
            f"Sessions in {session.status.value!r} state do not support imports."
        )


def _verify_concurrent_conflicts(
    preview: ParticipantSubmissionImportPreview,
    parsed_by_row: Mapping[int, _ParsedRow],
    current: Mapping[str, str],
) -> None:
    for planned in preview.rows:
        parsed = parsed_by_row[planned.row_number]
        if parsed.reference_digest is None:
            raise ParticipantSubmissionImportError("A ready row lost its reference binding.")
        actual = current.get(parsed.reference_digest)
        if actual != planned.existing_participant_id:
            raise ParticipantSubmissionImportError(
                "Participant conflicts changed after preview. Preview the file again."
            )


def _already_applied_result(
    batch: ParticipantSubmissionImportBatch,
) -> ApplyParticipantSubmissionImportResult:
    return ApplyParticipantSubmissionImportResult(
        batch_id=batch.batch_id,
        file_hash=batch.file_hash,
        plan_hash=batch.plan_hash,
        created_participant_count=0,
        enrolled_count=0,
        draft_count=0,
        submitted_count=0,
        skipped_count=batch.row_count,
        replaced_count=0,
        identity_count=0,
        imported_at=batch.imported_at,
        rows=(),
        result_report_csv=b"",
        already_applied=True,
    )


def _result_report(rows: Sequence[AppliedParticipantImportRow]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerow(("row_number", "participant_ref", "status", "message"))
    for row in rows:
        writer.writerow((row.row_number, _csv_safe(row.participant_ref), row.status, row.message))
    return stream.getvalue().encode("utf-8")


def _unreachable_group() -> SessionStakeholderGroup:
    raise AssertionError("A validated import row must contain a stakeholder group.")


@dataclass(frozen=True, slots=True)
class RedactExpiredImportedIdentityCommand:
    actor_id: str = "identity-retention-maintenance"
    limit: int = 100


@dataclass(frozen=True, slots=True)
class RedactExpiredImportedIdentityResult:
    redacted_count: int
    participant_ids: tuple[str, ...]


class RedactExpiredImportedIdentity:
    """Idempotently remove expired imported PII from durable storage."""

    def __init__(
        self,
        unit_of_work_factory: UnitOfWorkFactory,
        *,
        clock: Clock = utc_now,
        id_factory: IdFactory = lambda: str(new_id()),
    ) -> None:
        self._unit_of_work_factory = unit_of_work_factory
        self._clock = clock
        self._id_factory = id_factory

    def execute(
        self, command: RedactExpiredImportedIdentityCommand
    ) -> RedactExpiredImportedIdentityResult:
        if not command.actor_id.strip() or not 1 <= command.limit <= 500:
            raise ParticipantSubmissionImportError(
                "Identity maintenance requires an actor and a limit from 1 to 500."
            )
        at = self._clock()
        redacted_ids: list[str] = []
        with self._unit_of_work_factory() as unit_of_work:
            identities = unit_of_work.participant_identities.list_expired(
                at=at, limit=command.limit
            )
            for identity in identities:
                participant = unit_of_work.participants.get(identity.participant_id)
                if participant is None:
                    raise ParticipantSubmissionImportError(
                        "An expired identity has no analytical participant."
                    )
                redacted = identity.redact(actor_id=command.actor_id, at=at)
                unit_of_work.participant_identities.save(redacted)
                unit_of_work.audit_events.add(
                    operational_audit_event(
                        event_id=self._id_factory(),
                        occurred_at=at,
                        session_id=participant.session_id,
                        actor_id=command.actor_id,
                        actor_type=ActorType.SYSTEM,
                        action=AuditAction.UPDATED,
                        entity_type="participant_identity",
                        entity_id=participant.participant_id,
                        correlation_id=self._id_factory(),
                        use_case="redact_expired_imported_identity",
                        before_json={"identity_state": "encrypted"},
                        after_json={"identity_state": "redacted"},
                        reason_text="Imported identity retention period expired",
                    )
                )
                redacted_ids.append(participant.participant_id)
            unit_of_work.commit()
        return RedactExpiredImportedIdentityResult(
            redacted_count=len(redacted_ids), participant_ids=tuple(redacted_ids)
        )


@dataclass(frozen=True, slots=True)
class GenerateImportedResumeLinksCommand:
    session_id: str
    participant_ids: tuple[str, ...]
    actor_id: str
    actor_roles: frozenset[str]
    confirmed: bool
    expires_at: datetime | None = None
    correlation_id: str = field(default_factory=lambda: str(new_id()))

    def __post_init__(self) -> None:
        if "admin" not in self.actor_roles or not self.confirmed:
            raise ParticipantSubmissionImportError(
                "Resume-link generation requires an authorized administrator "
                "and explicit confirmation."
            )
        if not self.participant_ids or len(self.participant_ids) != len(
            set(self.participant_ids)
        ):
            raise ParticipantSubmissionImportError(
                "Select one or more unique imported participants."
            )


@dataclass(frozen=True, slots=True, repr=False)
class GeneratedImportedResumeCredential:
    participant_id: str
    participant_alias: str
    access_token: str
    expires_at: datetime


@dataclass(frozen=True, slots=True, repr=False)
class GenerateImportedResumeLinksResult:
    credentials: tuple[GeneratedImportedResumeCredential, ...]
    session_slug: str
    session_status: SessionStatus

    def __repr__(self) -> str:
        return (
            "GenerateImportedResumeLinksResult("
            f"credential_count={len(self.credentials)}, tokens=<redacted>)"
        )


class GenerateImportedResumeLinks:
    """Issue initial one-time resume credentials after import."""

    def __init__(
        self,
        unit_of_work_factory: UnitOfWorkFactory,
        *,
        clock: Clock = utc_now,
        id_factory: IdFactory = lambda: str(new_id()),
        grant_ttl: timedelta = timedelta(days=30),
    ) -> None:
        self._unit_of_work_factory = unit_of_work_factory
        self._clock = clock
        self._id_factory = id_factory
        self._grant_ttl = grant_ttl

    def execute(
        self, command: GenerateImportedResumeLinksCommand
    ) -> GenerateImportedResumeLinksResult:
        at = self._clock()
        expires_at = command.expires_at or at + self._grant_ttl
        if expires_at <= at:
            raise ParticipantSubmissionImportError(
                "Resume-link expiration must be in the future."
            )
        credentials: list[GeneratedImportedResumeCredential] = []
        with self._unit_of_work_factory() as unit_of_work:
            session = unit_of_work.session.get_for_update(command.session_id)
            if session is None or session.status in {
                SessionStatus.CANCELED,
                SessionStatus.ARCHIVED,
            }:
                raise ParticipantSubmissionImportError(
                    "This session cannot receive imported resume links."
                )
            for participant_id in command.participant_ids:
                participant = unit_of_work.participants.get_for_update(participant_id)
                if participant is None or participant.session_id != command.session_id:
                    raise ParticipantSubmissionImportError(
                        "A selected participant is unavailable."
                    )
                if unit_of_work.participant_imports.get_reference_for_participant(
                    command.session_id, participant_id
                ) is None:
                    raise ParticipantSubmissionImportError(
                        "Resume links may be generated here only for imported participants."
                    )
                active = unit_of_work.access_grants.get_current_for_participant_for_update(
                    participant_id, at=at
                )
                if active is not None:
                    raise ParticipantSubmissionImportError(
                        "A selected participant already has an active resume link."
                    )
                issued = generate_token()
                grant = ParticipantAccessGrant(
                    access_grant_id=self._id_factory(),
                    participant_id=participant_id,
                    token_digest=issued.token_digest,
                    issued_at=at,
                    expires_at=expires_at,
                )
                unit_of_work.access_grants.add(grant)
                unit_of_work.audit_events.add(
                    operational_audit_event(
                        event_id=self._id_factory(),
                        occurred_at=at,
                        session_id=command.session_id,
                        actor_id=command.actor_id,
                        actor_type=ActorType.USER,
                        action=AuditAction.CREATED,
                        entity_type="participant_access_grant",
                        entity_id=grant.access_grant_id,
                        correlation_id=command.correlation_id,
                        use_case="generate_imported_resume_links",
                        after_json={
                            "participant_id": participant_id,
                            "expires_at": expires_at,
                            "source": "administrator_import_follow_up",
                        },
                    )
                )
                credentials.append(GeneratedImportedResumeCredential(
                    participant_id=participant_id,
                    participant_alias=participant.alias or participant_id,
                    access_token=issued.plaintext,
                    expires_at=expires_at,
                ))
            unit_of_work.commit()
        return GenerateImportedResumeLinksResult(
            credentials=tuple(credentials),
            session_slug=session.public_slug,
            session_status=session.status,
        )
