from __future__ import annotations

import io
from datetime import UTC, datetime
from decimal import Decimal
from types import SimpleNamespace
from typing import Any, Self

import pytest
from cryptography.exceptions import InvalidTag
from openpyxl import load_workbook

from poli_insight.application.use_cases.participant_submission_import import (
    GenerateParticipantImportTemplate,
    GenerateParticipantImportTemplateCommand,
    ImportRowStatus,
    ParticipantSubmissionImportError,
    PreviewParticipantSubmissionImport,
    PreviewParticipantSubmissionImportCommand,
)
from poli_insight.domain.enum import (
    QuestionType,
    ResponseFormat,
    ResponseTargetType,
    SessionStatus,
)
from poli_insight.infrastructure.security.identity_protection import (
    AesGcmIdentityProtector,
    normalize_participant_reference,
)

NOW = datetime(2026, 8, 6, 16, tzinfo=UTC)


class _ImportRepository:
    def participants_by_reference_digests(
        self, session_id: str, digests: tuple[str, ...]
    ) -> dict[str, str]:
        del session_id, digests
        return {}


class _ParticipantRepository:
    def get_many(self, participant_ids: tuple[str, ...]) -> tuple[()]:
        del participant_ids
        return ()


class _UnitOfWork:
    def __init__(self, *, response_format: ResponseFormat = ResponseFormat.DIRECT_RATING) -> None:
        group = SimpleNamespace(
            session_stakeholder_group_id="group-id",
            configuration_version_id="configuration-id",
            group_key="community",
            name="Community",
            display_order=1,
            is_active=True,
        )
        question = SimpleNamespace(
            question_definition_id="question-id",
            configuration_version_id="configuration-id",
            question_key="criterion-cost",
            question_type=(
                QuestionType.ALTERNATIVE_RANK
                if response_format == ResponseFormat.DIRECT_RANKING
                else QuestionType.CRITERION_RATING
            ),
            display_order=1,
            required=True,
            prompt_snapshot="Rate the importance of Cost",
            criterion_id="criterion-id",
            left_criterion_id=None,
            right_criterion_id=None,
            alternative_id=(
                "alternative-id"
                if response_format == ResponseFormat.DIRECT_RANKING
                else None
            ),
        )
        configuration = SimpleNamespace(
            configuration_version_id="configuration-id",
            version_number=3,
            scenario_snapshot_id="snapshot-id",
            response_format=response_format,
            response_target_type=(
                ResponseTargetType.ALTERNATIVE
                if response_format == ResponseFormat.DIRECT_RANKING
                else ResponseTargetType.CRITERION
            ),
            scale_id="scale-id",
            allow_resubmissions=False,
            allow_incomplete_submission=False,
            consistency_threshold=Decimal("0.1"),
            config_hash="c" * 64,
            stakeholder_groups=(group,),
            question_definitions=(question,),
            is_activated=True,
            configuration_json={},
        )
        session = SimpleNamespace(
            session_id="session-id",
            public_slug="budget-session",
            title="Budget session",
            status=SessionStatus.DRAFT,
            identity_policy="pseudonymous",
            active_configuration=configuration,
        )
        scale_value = SimpleNamespace(
            scale_value_id="scale-value-id",
            stable_value_key="direct_1",
            label="Very Low",
            ordinal=1,
            numeric_value=Decimal(1),
        )
        snapshot = SimpleNamespace(
            scenario_snapshot_id="snapshot-id",
            title="Budget choices",
            declared_version="2026.1",
            root_hash="s" * 64,
            criteria=(
                SimpleNamespace(
                    criterion_id="criterion-id", criterion_key="cost"
                ),
            ),
            alternatives=(),
            scales=(
                SimpleNamespace(
                    scale_id="scale-id",
                    name="Five point",
                    definition_version=1,
                    values=(scale_value,),
                ),
            ),
        )
        self.session = SimpleNamespace(get=lambda session_id: session)
        self.scenarios = SimpleNamespace(get_by_id=lambda snapshot_id: snapshot)
        self.participant_imports = _ImportRepository()
        self.participants = _ParticipantRepository()

    def __enter__(self) -> Self:
        return self

    def __exit__(self, *args: object) -> None:
        del args


def _factory(response_format: ResponseFormat = ResponseFormat.DIRECT_RATING) -> Any:
    return lambda: _UnitOfWork(response_format=response_format)


def _protector() -> AesGcmIdentityProtector:
    return AesGcmIdentityProtector(
        hmac_secret="h" * 32,
        encryption_secret="e" * 32,
    )


def _pairwise_factory() -> Any:
    def factory() -> _UnitOfWork:
        unit = _UnitOfWork()
        session = unit.session.get("session-id")
        configuration = session.active_configuration
        question = configuration.question_definitions[0]
        configuration.response_format = ResponseFormat.PAIRWISE
        question.question_type = QuestionType.CRITERION_PAIR
        question.criterion_id = None
        question.left_criterion_id = "criterion-a-id"
        question.right_criterion_id = "criterion-b-id"
        snapshot = unit.scenarios.get_by_id("snapshot-id")
        snapshot.criteria = (
            SimpleNamespace(
                criterion_id="criterion-a-id", criterion_key="cost"
            ),
            SimpleNamespace(
                criterion_id="criterion-b-id", criterion_key="reliability"
            ),
        )
        return unit

    return factory


def _alternative_factory() -> Any:
    def factory() -> _UnitOfWork:
        unit = _UnitOfWork()
        session = unit.session.get("session-id")
        configuration = session.active_configuration
        question = configuration.question_definitions[0]
        configuration.response_target_type = ResponseTargetType.ALTERNATIVE
        question.question_type = QuestionType.ALTERNATIVE_RATING
        question.criterion_id = None
        question.alternative_id = "alternative-id"
        snapshot = unit.scenarios.get_by_id("snapshot-id")
        snapshot.alternatives = (
            SimpleNamespace(
                alternative_id="alternative-id", alternative_key="north"
            ),
        )
        return unit

    return factory


def test_templates_are_bound_deterministic_and_xlsx_has_dropdowns() -> None:
    generator = GenerateParticipantImportTemplate(_factory())
    csv_template = generator.execute(
        GenerateParticipantImportTemplateCommand(
            session_id="session-id", file_format="csv", filled_example=True
        )
    )
    assert csv_template.content.startswith(b"template_schema_version,session_slug")
    assert b"answer.criterion_rating.cost | Rate the importance of Cost" in csv_template.content
    assert b"budget-session" in csv_template.content

    excel_template = generator.execute(
        GenerateParticipantImportTemplateCommand(
            session_id="session-id", file_format="xlsx", filled_example=True
        )
    )
    workbook = load_workbook(io.BytesIO(excel_template.content))
    assert workbook.sheetnames == [
        "Responses",
        "Instructions",
        "Scale values",
        "Groups",
        "Configuration",
    ]
    assert len(workbook["Responses"].data_validations.dataValidation) >= 3
    assert workbook["Configuration"]["B1"].value == "1"
    assert workbook["Scale values"]["A2"].value == "direct_1"
    excel_preview = PreviewParticipantSubmissionImport(
        _factory(),
        _protector(),
        max_bytes=1_000_000,
        max_rows=100,
        clock=lambda: NOW,
    ).execute(
        PreviewParticipantSubmissionImportCommand(
            session_id="session-id",
            filename=excel_template.filename,
            content=excel_template.content,
        )
    )
    assert excel_preview.rows[0].status == ImportRowStatus.READY


def test_pairwise_orientation_and_alternative_targets_are_machine_readable() -> None:
    pairwise = GenerateParticipantImportTemplate(_pairwise_factory()).execute(
        GenerateParticipantImportTemplateCommand(
            session_id="session-id", file_format="csv"
        )
    )
    assert b"answer.criterion_pair.cost.reliability |" in pairwise.content
    assert b"answer.criterion_pair.reliability.cost |" not in pairwise.content

    alternative = GenerateParticipantImportTemplate(_alternative_factory()).execute(
        GenerateParticipantImportTemplateCommand(
            session_id="session-id", file_format="csv"
        )
    )
    assert b"answer.alternative_rating.north |" in alternative.content


def test_preview_accepts_generated_csv_and_redacts_representations() -> None:
    generator = GenerateParticipantImportTemplate(_factory())
    template = generator.execute(
        GenerateParticipantImportTemplateCommand(
            session_id="session-id", file_format="csv", filled_example=True
        )
    )
    previewer = PreviewParticipantSubmissionImport(
        _factory(), _protector(), max_bytes=1_000_000, max_rows=100,
        clock=lambda: NOW,
    )
    command = PreviewParticipantSubmissionImportCommand(
        session_id="session-id", filename="responses.csv", content=template.content
    )
    preview = previewer.execute(command)

    assert preview.summary.total_rows == 1
    assert preview.rows[0].status == ImportRowStatus.READY
    assert preview.rows[0].answer_count == 1
    assert "example-participant-001" not in repr(preview.rows[0])
    assert "template_schema_version" not in repr(command)


def test_formula_cells_and_unsupported_ranking_are_rejected() -> None:
    generator = GenerateParticipantImportTemplate(_factory())
    template = generator.execute(
        GenerateParticipantImportTemplateCommand(
            session_id="session-id", file_format="xlsx", filled_example=True
        )
    )
    workbook = load_workbook(io.BytesIO(template.content))
    workbook["Responses"]["E2"] = "=1+1"
    output = io.BytesIO()
    workbook.save(output)
    previewer = PreviewParticipantSubmissionImport(
        _factory(), _protector(), max_bytes=1_000_000, max_rows=100,
        clock=lambda: NOW,
    )
    with pytest.raises(ParticipantSubmissionImportError, match="Formula"):
        previewer.execute(
            PreviewParticipantSubmissionImportCommand(
                session_id="session-id",
                filename="responses.xlsx",
                content=output.getvalue(),
            )
        )
    with pytest.raises(ParticipantSubmissionImportError, match="not supported"):
        GenerateParticipantImportTemplate(
            _factory(ResponseFormat.DIRECT_RANKING)
        ).execute(
            GenerateParticipantImportTemplateCommand(
                session_id="session-id", file_format="csv"
            )
        )


def test_reference_digest_is_keyed_and_identity_encryption_is_authenticated() -> None:
    first = _protector()
    second = AesGcmIdentityProtector(
        hmac_secret="x" * 32,
        encryption_secret="e" * 32,
    )
    normalized = normalize_participant_reference("  Ａ-001  ")
    assert normalized == "A-001"
    assert first.participant_reference_digest(normalized) != second.participant_reference_digest(normalized)
    encrypted = first.encrypt("authored@example.test", context="participant:1")
    assert b"authored@example.test" not in encrypted
    assert first.decrypt(encrypted, context="participant:1") == "authored@example.test"
    with pytest.raises(InvalidTag):
        first.decrypt(encrypted, context="participant:2")
